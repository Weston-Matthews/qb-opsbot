# qb-opsbot
# Recording sub-module
# Written by Weston Matthews (2020)
# Takes a contract number, a year, and the settlement file path, then marks the contract as entered
# in an optional tracking workbook (one sheet per year) and links to the folder the settlement came from.
# The workbook path and its columns come from config.json; if no workbook is configured this step is skipped.

import os.path as path

from qb_opsbot import settings


def main(contract_num, year, contract_path, config=None):
    config = config or settings.load_config()
    wb_path = config.path("contract_summary_workbook")
    columns = config["record"]
    year = str(year)
    folder_path = contract_path[:max(contract_path.rfind('\\'), contract_path.rfind('/'))]

    if not wb_path or not path.exists(wb_path):
        # The spreadsheet doesn't exist; tell the settlement module to skip this step
        return 1

    import openpyxl as xl
    # The spreadsheet exists; find the right page, and record that the contract is done and where its folder is
    print("Recording successful task in summary spreadsheet...")
    wb_read = xl.load_workbook(wb_path, data_only=True)
    wb_write = xl.load_workbook(wb_path)
    if year not in wb_read.sheetnames:
        print("No sheet for " + year + " in recording spreadsheet. No big deal!")
        return 1
    ws_read = wb_read[year]
    ws_write = wb_write[year]

    # Find the contract number
    for row in range(1, ws_read.max_row + 1):
        if ws_read[columns["contract_col"] + str(row)].value == contract_num:
            ws_write[columns["done_col"] + str(row)] = 'x'
            link = ws_write[columns["link_col"] + str(row)]
            link.hyperlink = folder_path
            link.value = folder_path
            link.style = "Hyperlink"
            try:
                wb_write.save(wb_path)
                return 0
            except OSError:
                print("Error accessing spreadsheet...")
                return -1
    print("Did not find contract number in recording spreadsheet. No big deal!")
    return 1


if __name__ == "__main__":
    print("Sub-module to be used only by the settlement module!")
