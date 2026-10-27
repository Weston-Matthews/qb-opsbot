# Writes samples/sample_statement.xlsx: a fictional settlement statement on the third sheet,
# matching the layout in config.example.json, so the parser can be tried without real data.
#   python samples/make_sample_statement.py
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sample_statement import build_grid  # noqa: E402


def main():
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_statement.xlsx")
    wb = openpyxl.Workbook()
    wb.active.title = "Cover"
    wb.create_sheet("Lots")
    ws = wb.create_sheet("Statement")   # sheet index 2
    for r, row in enumerate(build_grid()):
        for c, value in enumerate(row):
            if value != "":
                ws.cell(row=r + 1, column=c + 1, value=value)
    wb.save(out)
    print("Wrote " + out)


if __name__ == "__main__":
    main()
